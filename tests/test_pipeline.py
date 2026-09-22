"""Offline end-to-end test of the pipeline: Gemini and Crustdata are faked,
so this exercises the orchestration (one drafting call per company,
search_name used for Crustdata, per-company checkpointing, QA flags)
without any network calls or cost."""
import os

from openpyxl import load_workbook

from src import pipeline
from src.config import Secrets
from src.crustdata_client import Contact
from src.discovery import Company


class FakeCrustdata:
    searched: list[str] = []

    def __init__(self, api_key):
        self.counter = type("C", (), {"summary": lambda self: "fake summary"})()

    def get_two_contacts(self, company_name):
        FakeCrustdata.searched.append(company_name)
        if company_name == "Nobody Corp":
            return Contact(role="hr_ta"), Contact(role="ops_scm")
        return (
            Contact(role="hr_ta", name="Priya Sharma", title="HR Manager",
                    linkedin_url="https://www.linkedin.com/in/priya"),
            Contact(role="ops_scm", name="Ravi Kumar", title=None,
                    linkedin_url="https://www.linkedin.com/in/ravi"),
        )

    def close(self):
        pass


def test_pipeline_one_draft_call_per_company(monkeypatch, tmp_path):
    FakeCrustdata.searched = []
    draft_calls: list[list[str]] = []

    monkeypatch.setattr(pipeline, "rank_companies", lambda sector, count, key: [
        Company(name="Cipla Limited", search_name="Cipla", rationale="large pharma"),
        Company(name="Nobody Corp Pvt Ltd", search_name="Nobody Corp", rationale="x"),
    ])
    monkeypatch.setattr(pipeline, "CrustdataClient", FakeCrustdata)

    def fake_draft(sector, company, contacts):
        from src.drafting import DraftedMessages
        found = [c for c in contacts if c.name]
        draft_calls.append([c.role for c in found])
        return {
            c.role: DraftedMessages(
                connection_note=f"Hi {c.name.split()[0]}, keen to connect on {sector} roles.",
                follow_up_message=f"Thanks for connecting, {c.name.split()[0]}.",
            )
            for c in found
        }

    monkeypatch.setattr(pipeline, "draft_company_messages", fake_draft)
    monkeypatch.setattr(pipeline, "build_output_path",
                        lambda sector, mode: str(tmp_path / "out.xlsx"))

    path = pipeline.run_pipeline("Pharma", 2, "test", Secrets("c", "g"))

    # Crustdata searched by the short search_name, not the legal name
    assert FakeCrustdata.searched == ["Cipla", "Nobody Corp"]
    # Exactly one drafting call per company, covering both roles together
    assert draft_calls == [["hr_ta", "ops_scm"], []]

    ws = load_workbook(path).active
    assert ws.max_row == 3  # header + 2 companies
    assert ws.cell(row=2, column=1).value == "Cipla Limited"  # display uses full name
    assert ws.cell(row=2, column=5).value.startswith("Hi Priya")
    assert ws.cell(row=2, column=10).value.startswith("Hi Ravi")
    assert not ws.cell(row=2, column=12).value  # clean row, no QA flag
    assert "no contact found" in ws.cell(row=3, column=12).value


def test_pipeline_drafting_failure_is_flagged_not_fatal(monkeypatch, tmp_path):
    monkeypatch.setattr(pipeline, "rank_companies", lambda sector, count, key: [
        Company(name="Cipla Limited", search_name="Cipla", rationale="large pharma"),
    ])
    monkeypatch.setattr(pipeline, "CrustdataClient", FakeCrustdata)

    def failing_draft(*args):
        raise RuntimeError("Gemini call failed on all models: 503")

    monkeypatch.setattr(pipeline, "draft_company_messages", failing_draft)
    monkeypatch.setattr(pipeline, "build_output_path",
                        lambda sector, mode: str(tmp_path / "out.xlsx"))

    path = pipeline.run_pipeline("Pharma", 1, "test", Secrets("c", "g"))

    ws = load_workbook(path).active
    # contacts still written even though drafting failed
    assert ws.cell(row=2, column=2).value == "Priya Sharma"
    assert "drafting failed" in ws.cell(row=2, column=12).value
    assert os.path.exists(path)
