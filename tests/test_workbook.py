import os
import tempfile

from openpyxl import load_workbook

from src.workbook import HEADERS, OutputRow, build_output_path, write_workbook


def _sample_row(**overrides) -> OutputRow:
    base = dict(
        company="Cipla Limited",
        hr_name="Priya Sharma",
        hr_title="Head of Talent Acquisition",
        hr_linkedin="https://www.linkedin.com/in/priya-example",
        hr_note="Hi Priya, ...",
        hr_followup="Thanks for connecting, Priya.",
        ops_name=None,
        ops_title=None,
        ops_linkedin=None,
        ops_note=None,
        ops_followup=None,
        qa_flag="",
    )
    base.update(overrides)
    return OutputRow(**base)


def test_build_output_path_test_mode_tagged():
    with tempfile.TemporaryDirectory() as tmp:
        path = build_output_path("Pharma", "test", output_dir=tmp)
        assert "_TEST_" in os.path.basename(path)


def test_build_output_path_full_mode_untagged():
    with tempfile.TemporaryDirectory() as tmp:
        path = build_output_path("Pharma", "full", output_dir=tmp)
        assert "_TEST" not in os.path.basename(path)


def test_writes_headers_and_row():
    with tempfile.TemporaryDirectory() as tmp:
        path = os.path.join(tmp, "out.xlsx")
        write_workbook([_sample_row()], path)
        assert os.path.exists(path)

        wb = load_workbook(path)
        ws = wb.active
        header_row = [c.value for c in ws[1]]
        assert header_row == HEADERS
        assert ws.cell(row=2, column=1).value == "Cipla Limited"


def test_linkedin_cell_is_real_hyperlink_not_bare_text():
    with tempfile.TemporaryDirectory() as tmp:
        path = os.path.join(tmp, "out.xlsx")
        write_workbook([_sample_row()], path)
        wb = load_workbook(path)
        ws = wb.active
        # HR/TA LinkedIn column is column 4
        cell = ws.cell(row=2, column=4)
        assert cell.hyperlink is not None
        assert cell.hyperlink.target == "https://www.linkedin.com/in/priya-example"


def test_missing_contact_leaves_cells_blank_not_broken_link():
    with tempfile.TemporaryDirectory() as tmp:
        path = os.path.join(tmp, "out.xlsx")
        write_workbook([_sample_row()], path)
        wb = load_workbook(path)
        ws = wb.active
        # Ops/SCM name column is column 7 -- no ops contact in this fixture
        cell = ws.cell(row=2, column=7)
        assert cell.hyperlink is None
        # openpyxl round-trips an empty string as None on reload -- both
        # mean "nothing written", which is what we're asserting here.
        assert cell.value in (None, "")


def test_repeated_write_overwrites_same_path_as_checkpoint():
    """The pipeline calls write_workbook repeatedly against the same path
    as a checkpoint after every company -- confirm a second, longer call
    replaces the first rather than appending or erroring."""
    with tempfile.TemporaryDirectory() as tmp:
        path = os.path.join(tmp, "out.xlsx")
        write_workbook([_sample_row()], path)
        write_workbook([_sample_row(), _sample_row(company="Sun Pharma")], path)

        wb = load_workbook(path)
        ws = wb.active
        assert ws.max_row == 3  # header + 2 data rows, not 1 + 1 + 2
        assert ws.cell(row=3, column=1).value == "Sun Pharma"
