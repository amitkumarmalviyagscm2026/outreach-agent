"""xlsx writer. LinkedIn URLs become real clickable cell hyperlinks -- never
bare text -- per the user's standing preference from interactive runs.

write_workbook() takes an explicit path rather than computing one itself,
so the pipeline can call it repeatedly against the SAME path as a
checkpoint after every company, not just once at the very end. A real
100-company run showed this matters: sustained Gemini 503s can push total
runtime well past a job timeout, and a workbook that only ever gets
written once, right before the function returns, loses 100% of a run's
work the moment it's killed -- even if 90 of 100 companies finished fine.
"""
from __future__ import annotations

import datetime as dt
import os
from dataclasses import dataclass

from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

HEADERS = [
    "Company",
    "HR/TA Name", "HR/TA Title", "HR/TA LinkedIn",
    "HR/TA Connection Note", "HR/TA Follow-up",
    "Ops/SCM Name", "Ops/SCM Title", "Ops/SCM LinkedIn",
    "Ops/SCM Connection Note", "Ops/SCM Follow-up",
    "QA Flag",
]

COLUMN_WIDTHS = [28, 20, 26, 34, 40, 40, 20, 26, 34, 40, 40, 40]


@dataclass
class OutputRow:
    company: str
    hr_name: str | None
    hr_title: str | None
    hr_linkedin: str | None
    hr_note: str | None
    hr_followup: str | None
    ops_name: str | None
    ops_title: str | None
    ops_linkedin: str | None
    ops_note: str | None
    ops_followup: str | None
    qa_flag: str


def build_output_path(sector: str, mode: str, output_dir: str = "output") -> str:
    """Computes the output file's path once, up front, so every checkpoint
    write during a run lands at the same path (overwriting in place)
    rather than each write picking a fresh timestamp and creating a new
    file. Call this once per run and reuse the result."""
    os.makedirs(output_dir, exist_ok=True)
    timestamp = dt.datetime.now().strftime("%Y%m%d_%H%M%S")
    safe_sector = "".join(c if c.isalnum() else "_" for c in sector)
    suffix = "_TEST" if mode == "test" else ""
    filename = f"outreach_{safe_sector}{suffix}_{timestamp}.xlsx"
    return os.path.join(output_dir, filename)


def _write_link_cell(ws: Worksheet, row: int, col: int, display_text: str | None, url: str | None) -> None:
    cell = ws.cell(row=row, column=col)
    if url:
        cell.value = display_text or url
        cell.hyperlink = url
        cell.style = "Hyperlink"
    else:
        cell.value = display_text or ""


def write_workbook(rows: list[OutputRow], path: str) -> str:
    """Writes (or overwrites) the workbook at `path`. Cheap enough --
    a handful of KB, sub-second -- to call after every company as a
    checkpoint, not just once at the end of a run."""
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Outreach"

    for col_idx, header in enumerate(HEADERS, start=1):
        ws.cell(row=1, column=col_idx, value=header)
    ws.freeze_panes = "A2"
    for col_idx, width in enumerate(COLUMN_WIDTHS, start=1):
        ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = width

    for r, row in enumerate(rows, start=2):
        ws.cell(row=r, column=1, value=row.company)
        _write_link_cell(ws, r, 2, row.hr_name, row.hr_linkedin)
        ws.cell(row=r, column=3, value=row.hr_title or "")
        _write_link_cell(ws, r, 4, row.hr_linkedin, row.hr_linkedin)
        ws.cell(row=r, column=5, value=row.hr_note or "")
        ws.cell(row=r, column=6, value=row.hr_followup or "")
        _write_link_cell(ws, r, 7, row.ops_name, row.ops_linkedin)
        ws.cell(row=r, column=8, value=row.ops_title or "")
        _write_link_cell(ws, r, 9, row.ops_linkedin, row.ops_linkedin)
        ws.cell(row=r, column=10, value=row.ops_note or "")
        ws.cell(row=r, column=11, value=row.ops_followup or "")
        ws.cell(row=r, column=12, value=row.qa_flag or "")

    wb.save(path)
    return path
